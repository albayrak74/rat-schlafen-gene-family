from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
def base():
    d=Document(); st=d.styles['Normal']; st.font.name='Times New Roman'; st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); st.font.size=Pt(11)
    s=d.sections[0]; s.page_height=Cm(29.7); s.page_width=Cm(21)
    for m in ['left_margin','right_margin','top_margin','bottom_margin']: setattr(s,m,Cm(2.2))
    return d
TITLE='The rat Schlafen gene family: genomic organization, phylogeny and expression across organs, aging and disease models'
P='../draft/Genetik_Makale_7_Gene_'
d=base(); d.add_paragraph().add_run('Highlights').bold=True
for l in open('highlights.txt').read().strip().split('\n'): d.add_paragraph(l,style='List Bullet')
d.save(P+'Highlights.docx')
# title page
d=base(); p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(TITLE); r.bold=True; r.font.size=Pt(14)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for t,s in [('İrem Alkan',0),('a',1),(', Mehmet Albayrak',0),('b,*',1)]: rr=p.add_run(t); rr.font.superscript=bool(s)
for l in ['ᵃ Department of Biology, Faculty of Arts and Sciences, Burdur Mehmet Akif Ersoy University, 15030 Burdur, Türkiye',
 'ᵇ Cloud Computing Program, Vocational School of Information Technologies, Isparta University of Applied Sciences, 32040 Isparta, Türkiye',
 '* Corresponding author: Mehmet Albayrak, Cloud Computing Program, Vocational School of Information Technologies, Isparta University of Applied Sciences, 32040 Isparta, Türkiye. E-mail: mehmetalbayrak@isparta.edu.tr',
 'E-mail addresses: iremyaman32@gmail.com (İ. Alkan), mehmetalbayrak@isparta.edu.tr (M. Albayrak)',
 'ORCID: İrem Alkan 0000-0003-1766-4586; Mehmet Albayrak 0000-0002-7089-122X']: d.add_paragraph(l)
import re; t=open('manuscript.md').read(); d.add_paragraph('Abbreviations: '+re.search(r'%ABBREV (.*)',t).group(1)).runs[0].font.size=Pt(10)
d.save(P+'TitlePage.docx')
# cover letter
d=base()
for txt in ['Dear Editor,',
f'We are pleased to submit our manuscript entitled “{TITLE}” for consideration as a research article in Gene.',
'Schlafen genes are interferon-inducible members of a rapidly evolving mammalian gene family whose proteins process transfer and ribosomal RNA, restrict viruses and influence immune-cell quiescence. The rat family has not been described systematically, and several of its gene symbols have changed between annotation releases, which makes earlier rat studies difficult to compare with current data. In this work we define the eight Schlafen genes of the current rat reference assembly (GRCr8), reconcile their names with earlier annotations, and describe their genomic organization, domain architecture, phylogeny with mouse and human, selective constraint and promoter motifs. We then examine their expression in eleven organs of rats at four ages and in six independent Wistar rat datasets on colitis, lung injury, adjuvant arthritis, IL-1β-exposed microglia, diet-induced obesity and liver aging.',
'The main findings are that the rat cluster shares its gene order with the mouse cluster but contains a single Slfn9 gene in place of mouse Slfn8 and Slfn9, and that the rat and mouse Slfn3 and Slfn4 genes are co-orthologues rather than one-to-one orthologues. The promoters are enriched for interferon-stimulated response elements. The broadly expressed members are more highly expressed in aged than in adult rats in several organs, and in kidney, lung, adrenal gland and liver these differences track a leukocyte marker. Expression in inflammation depends on the model, with large increases in inflamed colon and none after lipopolysaccharide lung injury. Cross-validated classifiers show that the age-group signal carried by Schlafen expression is modest and not stronger than that of random gene sets. We interpret all expression results as cross-sectional associations and discuss the limitations explicitly.',
'We believe the study fits the scope of Gene, which covers gene structure, gene family evolution and gene expression. The computational results are validated across independent datasets, organs and strains, and every analysis step is described in enough detail to be reproduced from public data.',
'This manuscript has not been published and is not under consideration for publication elsewhere. Both authors have approved the manuscript and agree with its submission. The authors declare no competing interests, and the study received no specific funding.',
'Thank you for considering our work.',
'Sincerely,','Mehmet Albayrak (corresponding author)','Cloud Computing Program, Vocational School of Information Technologies, Isparta University of Applied Sciences, 32040 Isparta, Türkiye','E-mail: mehmetalbayrak@isparta.edu.tr','on behalf of both authors (İrem Alkan, Mehmet Albayrak)']:
    d.add_paragraph(txt)
d.save(P+'CoverLetter.docx'); print('ok')
