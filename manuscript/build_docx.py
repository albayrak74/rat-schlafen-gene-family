import re,sys,os,pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__))); import render
MODE=sys.argv[1]; OUT=sys.argv[2]; LANG=sys.argv[3] if len(sys.argv)>3 else 'en'; TR=LANG=='tr'
t=open('manuscript_tr.md' if TR else 'manuscript.md').read()
title=re.search(r'%TITLE (.*)',t).group(1).strip()
abstract=t.split('%ABSTRACT')[1].split('%KEYWORDS')[0].strip()
kw=re.search(r'%KEYWORDS (.*)',t).group(1).strip(); abbr=re.search(r'%ABBREV (.*)',t).group(1).strip()
body=t.split(re.search(r'%ABBREV .*',t).group(0))[1]
caps=[c.strip() for c in open('captions_tr.md' if TR else 'captions.md').read().split('\n\n') if c.strip()]
d=Document(); st=d.styles['Normal']; st.font.name='Times New Roman'; st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); st.font.size=Pt(12)
pf=st.paragraph_format; pf.line_spacing=2.0 if MODE=='ms' else 1.15; pf.space_after=Pt(0 if MODE=='ms' else 6)
sec=d.sections[0]; sec.page_height=Cm(29.7); sec.page_width=Cm(21)
for m in ['left_margin','right_margin','top_margin','bottom_margin']: setattr(sec,m,Cm(2.5))
if MODE=='ms':
    ln=OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'),'1'); ln.set(qn('w:restart'),'continuous'); sec._sectPr.append(ln)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=fp.add_run(); f1=OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'),'begin'); it=OxmlElement('w:instrText'); it.text='PAGE'; f2=OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'),'end'); r._r.append(f1); r._r.append(it); r._r.append(f2)
GENES=r'(?<![\w-])(Slfnl1|Slfn1[34]|Slfn[1-9]|Irf7|Isg15|Mx1|Rsad2|Stat1)(?![\w-])'
def trcite(t):
    return re.sub(r'\(([^()]*\d{4}[a-f]?)\)',lambda m:'('+m.group(1).replace(' et al.,',' ve ark.,').replace(' and ',' ve ')+')',t) if TR else t
def para(text,bold=False,italic=False,size=None,align=None,space_after=None,genes=True):
    p=d.add_paragraph()
    if align is not None: p.alignment=align
    if genes: text=re.sub(GENES,r'*\1*',text)
    text=text.replace('Rattus norvegicus','*Rattus norvegicus*').replace('Helicobacter','*Helicobacter*').replace('in vitro','*in vitro*')
    text=re.sub(r'\*\*+','*',text)
    for s in re.split(r'(\*[^*]+\*)',text):
        if s.startswith('*') and s.endswith('*') and len(s)>2: rr=p.add_run(s[1:-1]); rr.italic=True
        else: rr=p.add_run(s); rr.italic=italic
        rr.bold=bold
        if size: rr.font.size=Pt(size)
    if space_after is not None: p.paragraph_format.space_after=Pt(space_after)
    return p
def table(df,widths,fs=8,italcol=None):
    tb=d.add_table(rows=1,cols=len(df.columns)); tb.style='Table Grid'
    for i,h in enumerate(df.columns): c=tb.rows[0].cells[i]; c.text=''; rr=c.paragraphs[0].add_run(h); rr.bold=True; rr.font.size=Pt(fs)
    for _,row in df.iterrows():
        cells=tb.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=''; rr=cells[i].paragraphs[0].add_run(str(v)); rr.font.size=Pt(fs); rr.italic=(italcol==i)
            cells[i].paragraphs[0].paragraph_format.line_spacing=1.0
    tb.autofit=False
    for i,col in enumerate(tb.columns): col.width=Cm(widths[i])
    for row in tb.rows:
        for i,c in enumerate(row.cells): c.width=Cm(widths[i])
    para('',genes=False)
T1=pd.read_csv('../out/Table1_datasets.csv'); T2=pd.read_csv('../out/Table2_genes.csv',dtype=str)
if TR:
    T1.columns=['Erişim no','Soy','Doku','Karşılaştırma','Kullanılan örnek','Veri tipi','Özgün çalışma']
    T2.columns=['Gen','Grup','GeneID','Konum (kr:başlangıç–bitiş, zincir)','Ekzon','Protein','Uzunluk (aa)','MA (kDa)','pI','Iİ','GRAVY','Aİ','Bölgeler']
CAP1=('Tablo 1 Bu çalışmada analiz edilen açık sıçan transkriptom veri setleri (Gene Expression Omnibus). Deneysel tedavi uygulanan gruplar analize alınmamıştır. FPKM, haritalanan milyon fragman başına transkript kilobazı başına fragman; IL-1β, interlökin-1β; LPS, lipopolisakkarit; PBS, fosfat tamponlu salin; TNBS, 2,4,6-trinitrobenzen sülfonik asit' if TR else 'Table 1 Public rat transcriptome datasets analyzed in this study (Gene Expression Omnibus). Groups that received experimental treatments were not analyzed. FPKM, fragments per kilobase of transcript per million mapped fragments; IL-1β, interleukin-1β; LPS, lipopolysaccharide; PBS, phosphate-buffered saline; TNBS, 2,4,6-trinitrobenzene sulfonic acid')
CAP2=('Tablo 2 Sıçan GRCr8 genomundaki Schlafen genleri ve kodladıkları proteinler (RefSeq GCF_036323735.1). MA, molekül ağırlığı; pI, izoelektrik nokta; Iİ, kararsızlık indeksi; GRAVY, ortalama hidropati; Aİ, alifatik indeks. Bölgeler (Pfam, toplama eşiği): N, SLFN N-terminal (PF17057); A, AlbA_2 (PF04326); G, SLFN GTPaz-benzeri (PF21026); H, SLFN grup 3 helikaz (PF09848); (H), yalnızca aileye özgü profille saptanan helikaz benzeri bölge. Konum ve ekzon sayısı temsilci transkripte aittir' if TR else 'Table 2 Schlafen genes of the rat GRCr8 assembly and their encoded proteins (RefSeq GCF_036323735.1). MW, molecular weight; pI, isoelectric point; II, instability index; GRAVY, grand average of hydropathy; AI, aliphatic index. Domains (Pfam, gathering thresholds): N, SLFN N-terminal domain (PF17057); A, AlbA_2 (PF04326); G, SLFN GTPase-like domain (PF21026); H, SLFN group 3 helicase domain (PF09848); (H), helicase-like region detected only with a family-specific profile. Location and exon number refer to the representative transcript')
def add_t1(): para(CAP1,size=10,genes=False); table(T1,[1.8,1.8,1.9,3.9,2.8,1.6,2.2],fs=7.5)
def add_t2(): para(CAP2,size=10); table(T2,[1.1,1.0,1.2,2.6,0.9,2.0,1.0,1.0,0.95,0.95,1.1,0.95,1.25],fs=6.5,italcol=0)
para(title,bold=True,size=14,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=12,genes=False)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for txt,sup in [('İrem Alkan',None),('a',True),(', Mehmet Albayrak',None),('b,*',True)]:
    rr=p.add_run(txt); rr.font.superscript=bool(sup)
for line in ['ᵃ Department of Biology, Faculty of Arts and Sciences, Burdur Mehmet Akif Ersoy University, 15030 Burdur, Türkiye',
             'ᵇ Cloud Computing Program, Vocational School of Information Technologies, Isparta University of Applied Sciences, 32040 Isparta, Türkiye',
             ('* Sorumlu yazar' if TR else '* Corresponding author')+': Mehmet Albayrak, Cloud Computing Program, Vocational School of Information Technologies, Isparta University of Applied Sciences, 32040 Isparta, Türkiye. E-mail: mehmetalbayrak@isparta.edu.tr',
             'ORCID: İrem Alkan 0000-0003-1766-4586; Mehmet Albayrak 0000-0002-7089-122X']:
    para(line,size=11,genes=False)
para(('Kısaltmalar: ' if TR else 'Abbreviations: ')+abbr,size=10,genes=False)
if MODE=='ms': d.add_page_break()
para('Özet' if TR else 'Abstract',bold=True); para(trcite(render.render(abstract)))
para(('Anahtar kelimeler: ' if TR else 'Keywords: ')+kw,space_after=12,genes=False)
if MODE=='ms': d.add_page_break()
figdone=set(); tdone=set()
for block in body.split('\n\n'):
    b=block.strip()
    if not b: continue
    if b.startswith('# '): h=d.add_paragraph(); rr=h.add_run(b[2:]); rr.bold=True; h.paragraph_format.space_before=Pt(12); continue
    if b.startswith('## '): h=d.add_paragraph(); rr=h.add_run(b[3:]); rr.italic=True; rr.bold=True; h.paragraph_format.space_before=Pt(6); continue
    para(trcite(render.render(b)))
    if MODE=='preview':
        for n in re.findall(r'Tablo (\d)' if TR else r'Table (\d)',b):
            if n not in tdone: tdone.add(n); (add_t1 if n=='1' else add_t2)()
        for n in re.findall(r'Şekil (\d)' if TR else r'Fig\. (\d)',b):
            if n not in figdone:
                figdone.add(n); d.add_picture(f'../fig/Fig{n}.png',width=Cm(16)); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
                para(caps[int(n)-1],size=10,space_after=10)
exec(open('decl.py').read())
for h,x in (DECL_TR if TR else DECL_EN): para(h,bold=True,genes=False); para(x,genes=False)
if MODE=='ms': d.add_page_break()
para('Kaynaklar' if TR else 'References',bold=True)
for r in render.reflist():
    p=para(r,size=11 if MODE=='ms' else 10,genes=False); p.paragraph_format.left_indent=Cm(0.75); p.paragraph_format.first_line_indent=Cm(-0.75)
if MODE=='ms':
    d.add_page_break(); add_t1(); add_t2()
    d.add_page_break(); para('Şekil altyazıları' if TR else 'Figure captions',bold=True)
    for c in caps: para(c)
d.save(OUT); print('saved',OUT)
