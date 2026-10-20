import pandas as pd, numpy as np, json, gzip, re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
TITLE='Comparative genomic and cross-dataset transcriptomic analysis of the rat Schlafen gene family in aging and inflammation'
o='../out/'
def base():
    d=Document(); st=d.styles['Normal']; st.font.name='Times New Roman'; st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); st.font.size=Pt(11)
    s=d.sections[0]; s.page_height=Cm(29.7); s.page_width=Cm(21)
    for m in ['left_margin','right_margin','top_margin','bottom_margin']: setattr(s,m,Cm(2.0))
    return d
d=base(); r=d.add_paragraph().add_run('Supplementary material'); r.bold=True; r.font.size=Pt(13)
d.add_paragraph(TITLE).runs[0].italic=True
d.add_paragraph('İrem Alkan, Mehmet Albayrak. Corresponding author: Mehmet Albayrak, Isparta University of Applied Sciences, mehmetalbayrak@isparta.edu.tr')
def tab(title,df,fs=7.5,note=None):
    d.add_paragraph().add_run(title).bold=True
    t=d.add_table(rows=1,cols=len(df.columns)); t.style='Table Grid'
    for i,c in enumerate(df.columns): cl=t.rows[0].cells[i]; cl.text=''; rr=cl.paragraphs[0].add_run(str(c)); rr.bold=True; rr.font.size=Pt(fs)
    for _,row in df.iterrows():
        cells=t.add_row().cells
        for i,v in enumerate(row):
            if isinstance(v,float): s='' if np.isnan(v) else (f'{v:.2e}' if (abs(v)<0.001 and v!=0) else f'{v:.3g}')
            else: s=str(v)
            cells[i].text=''; rr=cells[i].paragraphs[0].add_run(s); rr.font.size=Pt(fs)
    if note: d.add_paragraph(note).runs[0].font.size=Pt(8)
    d.add_paragraph('')
sheets={}
# S1
P=pd.read_csv(o+'slfn_rat_domains_pfam.csv'); P['search']='Pfam (gathering threshold)'
F=pd.read_csv(o+'slfn_rat_domains.csv'); F['search']='Family-specific profile (E < 1e-3)'; F['score']=np.nan
S1=pd.concat([P,F])[['gene','search','domain','start','end','score','ievalue']].sort_values(['gene','search','start'])
S1.columns=['Gene','Search','Domain','Start','End','Bit score','i-Evalue']
H=pd.read_csv(o+'proteome_pfam_hits.csv'); 
tab('Table S1. Schlafen domains in the eight rat proteins detected with Pfam models (gathering thresholds) and with family-specific profiles built from mouse and human Schlafen proteins. A proteome-wide Pfam search of all 97,483 RefSeq GRCr8 proteins detected B3R, AlbA_2, GTPase-like or group 3 helicase domains in 46 proteins, all isoforms of the eight genes listed (Slfn9, 28 isoforms; Slfnl1, 4; Slfn3, 4; Slfn14, 3; Slfn2, Slfn4 and Slfn5, 2 each; Slfn1, 1). A tblastn search of chromosome 10: 68.2–68.8 Mb with all mouse and human Schlafen proteins found one additional segment outside the gene models (10:68,562,629–68,563,220; similar to residues 393–597 of mouse Slfn4, 56% identity; E = 8e-56).',S1)
sheets['S1_domains']=S1; sheets['S1b_proteome_hits']=H
# S2
M=pd.read_csv(o+'slfn_promoter_motifs.csv'); E=pd.read_csv(o+'promoter_enrichment.csv'); E['P_Bonferroni']=np.minimum(E.P_resampling*4,1)
M.columns=['Gene','ISRE','GAS','NF-κB','IRF-E','GC (%)']
tab('Table S2. Consensus motif counts (both strands) in the 2 kb upstream of the transcription start site of each representative transcript, and enrichment relative to 200 random 2-kb windows of chromosome 10',M)
E.columns=['Motif','Observed (8 promoters)','Expected','P (resampling)','P (Poisson)','Promoters with motif','P (Bonferroni)']; tab('',E)
sheets['S2_motifs']=M; sheets['S2b_enrichment']=E
# S3
K=pd.read_csv(o+'slfn_kaks.csv'); K['type']=K['type'].replace({'orthologue':'rat–mouse pair at corresponding position','paralogue':'rat paralogues','cross pair':'rat–mouse cross pair'})
K.loc[K.Ks<0,['Ks','Ka_Ks']]=np.nan; K.columns=['Comparison','Gene 1','Gene 2','Ka','Ks','Ka/Ks']
tab('Table S3. Non-synonymous (Ka) and synonymous (Ks) substitution rates (Nei–Gojobori). Ks > 1 indicates saturation; an empty Ks indicates that it could not be estimated',K)
sheets['S3_KaKs']=K
# S4
B=pd.read_csv(o+'bodymap_slfn_age_sex.csv'); B=B[['test','organ','gene','baseMean','log2FC','lfcSE','p','padj_tx','fdr_family']]
B.columns=['Comparison','Organ','Gene','Mean count','log2FC','SE','P','FDR (transcriptome)','FDR (Schlafen family)']
tab('Table S4a. Age (104 vs 21 weeks; adjusted for sex) and sex (female vs male; adjusted for age) differences in Schlafen expression in the rat BodyMap (PyDESeq2)',B,fs=7)
A=pd.read_csv(o+'age_ptprc_adjusted.csv'); A.columns=['Organ','Gene','Age effect','P','Age effect (adjusted for Ptprc)','P (adjusted)']
tab('Table S4b. Age effect (aged vs adult) on log2(CPM + 1) estimated by linear regression within each organ, with sex as covariate where applicable, without and with adjustment for Ptprc expression',A,fs=7)
sheets['S4_BodyMap_DE']=B; sheets['S4b_Ptprc_adjusted']=A
# S5
W=pd.read_csv(o+'wistar_slfn_DE.csv'); W.columns=['Dataset','Comparison','Gene','Mean expression','log2FC','SE','P','FDR (transcriptome)','FDR (Schlafen family)']
tab('Table S5. Schlafen expression in the Wistar rat datasets. PyDESeq2 for count data; Welch t-test on log2(FPKM + 1) for GSE228906; linear model on normalized values for GSE160153 (transcriptome-wide FDR not computed for these two)',W,fs=7)
sheets['S5_Wistar_DE']=W
# S6
L=pd.read_csv(o+'age_ml_lolo.csv'); N=pd.read_csv(o+'age_ml_null.csv'); R=json.load(open(o+'revision_ml.json')); Sh=pd.read_csv(o+'age_ml_shap.csv',index_col=0).reset_index()
L['feature_set']=L.feature_set.str.replace('Transcriptome (2,000 most variable)','Transcriptome (2,000 genes, highest within-organ variance)',regex=False)
tab('Table S6a. Aged vs adult classification, leave-one-organ-out AUC (pooled and per held-out organ). en, elastic-net logistic regression; rf, random forest',L,fs=6.5)
sm=pd.DataFrame([['All organs','0.707 (95% CI 0.629–0.791)',round(N.pooled_AUC.median(),3),round(N.pooled_AUC.quantile(0.95),3),15,0.080],['Without testis',round(R['age_noTestis_SLFN'],3),round(R['age_noTestis_null_median'],3),'',int(round(R['age_noTestis_P']*201-1)),round(R['age_noTestis_P'],3)]],columns=['Analysis','Schlafen AUC','Random-set median','Random-set 95th percentile','Random sets ≥ observed (of 200)','P'])
tab('Table S6b. Comparison of the Schlafen elastic-net model with 200 random four-gene sets; P = (k + 1)/201',sm)
Sh.columns=['Gene','Mean |SHAP|','Direction (correlation of expression with SHAP)']; tab('Table S6c. SHAP importance in the random-forest age model (in-sample)',Sh)
sheets['S6a_age_LOLO']=L; sheets['S6b_age_null']=sm; sheets['S6c_age_SHAP']=Sh
# S7
I=pd.read_csv(o+'ml_lodo.csv'); I['feature_set']=I.feature_set.str.replace('Transcriptome (2,000 most variable)','Transcriptome (2,000 genes, highest within-dataset variance)',regex=False)
tab('Table S7a. Inflamed vs non-inflamed classification, leave-one-dataset-out AUC',I,fs=6.5)
U=pd.read_csv(o+'ml_univariate.csv'); U.columns=['Dataset','Gene','Within-dataset AUC (single gene, signed)']
tab('Table S7b. Within-dataset AUC of each Schlafen gene used alone (higher expression in inflamed samples gives AUC > 0.5)',U)
Si=pd.read_csv(o+'ml_shap.csv',index_col=0).reset_index(); Si.columns=['Gene','Mean |SHAP|','Direction']; tab('Table S7c. SHAP importance in the random-forest inflammation model (in-sample)',Si)
sheets['S7a_infl_LODO']=I; sheets['S7b_univariate']=U; sheets['S7c_infl_SHAP']=Si
out=sys.argv[1] if len(sys.argv)>1 else '../draft/Supplementary.docx'
d.save(out)
with pd.ExcelWriter(out.replace('.docx','_Data.xlsx')) as w:
    for n,df in sheets.items(): df.to_excel(w,sheet_name=n[:31],index=False)
print('ok')
